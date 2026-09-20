"""
    Copyright (C) 2025-26 Dipl.-Ing. Christoph Massmann <chris@dev-investor.de>

    This file is part of pp-terminal.
"""
# pylint: disable=too-many-locals
import re
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from pp_terminal.data.broker_import.models import ImportCandidate


def _excel_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).replace(second=0, microsecond=0)
    return datetime.fromisoformat(str(value)).replace(second=0, microsecond=0)


def _parse_trade_shares(comment: str) -> Decimal:
    match = re.search(r'\b(?:OPEN|CLOSE)\s+BUY\s+([0-9]+(?:\.[0-9]+)?)(?:/[0-9]+(?:\.[0-9]+)?)?\s+@', comment)
    if match is None:
        return Decimal('0')
    return Decimal(match.group(1))


def parse_xlsx(path: Path) -> list[ImportCandidate]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if 'Cash Operations' not in workbook.sheetnames:
        return []

    sheet = workbook['Cash Operations']
    sheet.reset_dimensions()
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next(index for index, row in enumerate(rows) if row and row[0] == 'Type')
    headers = [str(value) if value is not None else '' for value in rows[header_index]]
    raw_operations = [dict(zip(headers, row)) for row in rows[header_index + 1:] if row and row[0] not in (None, 'Total')]
    candidates: list[ImportCandidate] = []
    interest_tax_by_note: dict[str, Decimal] = {}
    withholding_tax_by_day_ticker: dict[tuple[object, str | None], Decimal] = {}
    dividends: dict[tuple[object, str | None, str], tuple[datetime, Decimal, str]] = {}

    for row in raw_operations:
        operation_type = str(row.get('Type') or '')
        comment = str(row.get('Comment') or '')
        ticker = str(row.get('Ticker') or '') or None
        amount = Decimal(str(row.get('Amount') or '0'))
        date = _excel_datetime(row.get('Time'))
        if operation_type == 'Free funds interest tax':
            interest_tax_by_note[comment.replace(' Tax', '')] = abs(amount)
        elif operation_type == 'Withholding tax':
            tax_key = (date.date(), ticker)
            withholding_tax_by_day_ticker[tax_key] = withholding_tax_by_day_ticker.get(tax_key, Decimal('0')) + abs(amount)
        elif operation_type == 'Dividend' and not comment.lower().startswith('corr '):
            dividend_key = (date.date(), ticker, comment)
            _, current_amount, external_id = dividends.get(dividend_key, (date, Decimal('0'), str(row.get('ID') or '')))
            dividends[dividend_key] = (date, current_amount + abs(amount), external_id or str(row.get('ID') or ''))
        elif operation_type in {'Stock purchase', 'Stock sell'}:
            shares = _parse_trade_shares(comment)
            transaction_type = 'BUY' if operation_type == 'Stock purchase' else 'SELL'
            external_id = str(row.get('ID') or f'{date.isoformat()}-{operation_type}-{ticker or ""}')
            candidates.append(ImportCandidate(external_id, date, transaction_type, ticker, None, 'EUR', shares, abs(amount), Decimal('0'), Decimal('0'), amount, comment))

    for (day, ticker, comment), (date, dividend_amount, external_id) in dividends.items():
        tax = withholding_tax_by_day_ticker.get((day, ticker), Decimal('0'))
        net_amount = dividend_amount - tax
        candidates.append(ImportCandidate(external_id, date, 'DIVIDENDS', ticker, None, 'EUR', Decimal('0'), net_amount, Decimal('0'), tax, net_amount, comment))

    for row in raw_operations:
        operation_type = str(row.get('Type') or '')
        if operation_type != 'Free funds interest':
            continue
        comment = str(row.get('Comment') or '')
        amount = Decimal(str(row.get('Amount') or '0'))
        tax = interest_tax_by_note.get(comment, Decimal('0'))
        net_amount = amount - tax
        date = _excel_datetime(row.get('Time'))
        external_id = str(row.get('ID') or f'{date.isoformat()}-{operation_type}')
        candidates.append(ImportCandidate(external_id, date, 'INTEREST', None, None, 'EUR', Decimal('0'), net_amount, Decimal('0'), tax, net_amount, comment))

    return candidates
